"""ForecastIQ: see next month's sales before it happens."""
from __future__ import annotations

import io
import json
import logging
import re
import secrets
import threading
import time
from datetime import date, datetime

import pandas as pd
from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from . import calendar, data, db, model, reports, service
from .config import (ADMIN_EMAIL, ADMIN_PASSWORD, APP_NAME, BRAND_DIR, CONTACT_EMAIL, CONTACT_WHATSAPP, DEMO_EMAIL, DEMO_MODE, DEMO_PASSWORD,
                     MAX_UPLOAD_MB, SAMPLE_DIR, STATIC_DIR)

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
log = logging.getLogger("forecast")
app = FastAPI(title=APP_NAME, docs_url=None, redoc_url=None)
COOKIE = "fq_session"
STATE = {"ready": False, "status": "starting"}


@app.on_event("startup")
def startup():
    db.init_schema()
    threading.Thread(target=_boot, daemon=True).start()


def _boot():
    try:
        if not db.one("SELECT id FROM events LIMIT 1"):
            evs = calendar.default_events() + (calendar.DEMO_PROMOTIONS if DEMO_MODE else [])
            for e in evs:
                db.execute("INSERT INTO events(name, start, end, kind, discount, source) VALUES(?,?,?,?,?,?)",
                           (e["name"], e["start"].isoformat(), e["end"].isoformat(), e["kind"], e["discount"], "default"))
        if DEMO_MODE and not db.one("SELECT id FROM users WHERE email=?", (DEMO_EMAIL,)):
            for email, name, role, pw in [(DEMO_EMAIL, "Demo Owner", "admin", DEMO_PASSWORD), ("finance@saharamart.example", "Nadia Hussain", "analyst", None),
                                          ("branches@saharamart.example", "Usman Tariq", "viewer", None)]:
                db.execute("INSERT INTO users(email, name, role, password_hash, created_at) VALUES(?,?,?,?,?)", (email, name, role, db.hash_password(pw) if pw else None, db.now()))
        if not DEMO_MODE and ADMIN_EMAIL and ADMIN_PASSWORD and not db.one("SELECT id FROM users LIMIT 1"):
            db.execute("INSERT INTO users(email, name, role, password_hash, created_at) VALUES(?,?,?,?,?)", (ADMIN_EMAIL, "Admin", "admin", db.hash_password(ADMIN_PASSWORD), db.now()))
        if DEMO_MODE and not db.one("SELECT id FROM datasets LIMIT 1"):
            STATE["status"] = "loading the demo company"
            did = service.save_dataset("Sahara Mart: daily sales (sample data)", "sahara_mart_sales.csv", (SAMPLE_DIR / "sahara_mart_sales.csv").read_bytes(),
                                       (SAMPLE_DIR / "sahara_mart_stock.csv").read_bytes(), is_demo=1, decisions={"missing": "missing"}, user="Demo Owner")
            db.execute("UPDATE datasets SET active=1 WHERE id=?", (did,))
        if service.active_id():
            STATE["status"] = "preparing forecasts"
            service.run()
            STATE["ready"] = True
            service.warm()
            if not db.one("SELECT id FROM reports LIMIT 1"):
                rid = reports.build()
                db.execute("UPDATE reports SET sent_to='demo (not sent)' WHERE id=?", (rid,)) if DEMO_MODE else None
        STATE["ready"] = True
        STATE["status"] = "ready"
    except Exception:
        log.exception("startup failed")
        STATE["status"] = "error"
        STATE["ready"] = True
    _scheduler()


def _scheduler():
    """On the report day each month, build and email the monthly forecast report."""
    while True:
        try:
            s = db.get_settings()
            today = date.today()
            if today.day == int(s.get("report_day", 1)) and service.active_id():
                month = (pd.Timestamp(service.load()["daily"]["date"].max()) + pd.offsets.MonthBegin(1)).strftime("%Y-%m")
                r = db.one("SELECT * FROM reports WHERE month=?", (month,))
                if not r or not r["sent_to"]:
                    reports.send(reports.build())
        except Exception:
            log.exception("scheduler error")
        time.sleep(3600)


# --------------------------------------------------------------------------- auth
def current_user(request: Request):
    token = request.cookies.get(COOKIE)
    u = db.one("SELECT u.id, u.email, u.name, u.role FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.token=?", (token,)) if token else None
    if not u:
        raise HTTPException(401, "Please sign in")
    return u


def need(*roles):
    def dep(user=Depends(current_user)):
        if user["role"] not in roles:
            raise HTTPException(403, "You don't have permission for this. Ask an admin.")
        return user
    return dep


EDIT = need("admin", "analyst")


def needs_data():
    if not service.active_id():
        raise HTTPException(409, "No sales data yet. Upload your sales file on the Data page.")
    if not STATE["ready"]:
        raise HTTPException(503, "Forecasts are being prepared. One moment…")


def _login(response, uid):
    token = secrets.token_urlsafe(32)
    db.execute("INSERT INTO sessions(token, user_id, created_at) VALUES(?,?,?)", (token, uid, db.now()))
    response.set_cookie(COOKIE, token, httponly=True, samesite="lax", max_age=60 * 60 * 24 * 14)


class LoginIn(BaseModel):
    email: str
    password: str


@app.post("/api/login")
def login(body: LoginIn, response: Response):
    u = db.one("SELECT * FROM users WHERE lower(email)=lower(?)", (body.email.strip(),))
    if not u or not db.check_password(body.password, u["password_hash"]):
        raise HTTPException(401, "Email or password is not correct")
    _login(response, u["id"])
    return {"ok": True}


@app.post("/api/demo-login")
def demo_login(response: Response):
    u = db.one("SELECT id FROM users WHERE email=?", (DEMO_EMAIL,))
    if not u or not STATE["ready"]:
        raise HTTPException(503, "The demo company is still loading. Try again in a few seconds.")
    _login(response, u["id"])
    return {"ok": True}


@app.post("/api/logout")
def logout(request: Request, response: Response):
    db.execute("DELETE FROM sessions WHERE token=?", (request.cookies.get(COOKIE, ""),))
    response.delete_cookie(COOKIE)
    return {"ok": True}


@app.get("/api/me")
def me(user=Depends(current_user)):
    return user


@app.get("/api/app-info")
def app_info():
    s = db.get_settings()
    did = service.active_id()
    d = db.one("SELECT name, is_demo FROM datasets WHERE id=?", (did,)) if did else None
    return {"app_name": APP_NAME, "company_name": s["company_name"], "accent_color": s["accent_color"], "logo_url": s["logo_url"],
            "currency": s["currency"], "ready": STATE["ready"], "status": STATE["status"], "dataset": d, "contact_email": CONTACT_EMAIL, "contact_whatsapp": CONTACT_WHATSAPP}


# --------------------------------------------------------------------------- screens
@app.get("/api/overview")
def overview(user=Depends(current_user)):
    needs_data()
    return service.overview()


def _filters(product, category, branch):
    return {"product": product or None, "category": category or None, "branch": branch or None}


@app.get("/api/explorer")
def explorer(product: str = "", category: str = "", branch: str = "", horizon: int = 3, view: str = "weekly", history: int = 12, user=Depends(current_user)):
    needs_data()
    if horizon not in (1, 3, 6, 12):
        horizon = 3
    try:
        r = service.run(**_filters(product, category, branch))
    except ValueError as e:
        raise HTTPException(404, str(e))
    s, fc, bt = r["series"], r["forecast"], r["backtest"]
    end = s.y.index.max()
    fend = pd.Timestamp(fc["monthly"][horizon - 1]["month"] + "-01") + pd.offsets.MonthEnd(0)
    hstart = s.y.index.min() if history == 0 else max(s.y.index.min(), end - pd.DateOffset(months=history))
    rule = {"daily": "D", "weekly": "W-SUN", "monthly": "MS"}.get(view, "W-SUN")
    if view == "monthly":
        hstart = hstart.replace(day=1)
    hist, fore = service.periods(s, fc, hstart, fend, rule, bt["range_pct"])
    if view == "monthly":
        mm = {m["month"]: m for m in fc["monthly"]}
        fore = [(t, v, mm[t.strftime("%Y-%m")]["low"], mm[t.strftime("%Y-%m")]["high"]) if t.strftime("%Y-%m") in mm else (t, v, l, h) for t, v, l, h in fore]
    excluded = s.exclude[s.exclude.index >= hstart]
    fmt = "%Y-%m-%d"
    evs = [{"date": str(e["start"]), "end": str(e["end"]), "name": e["name"], "kind": e["kind"]} for e in r["events"]
           if hstart <= pd.Timestamp(e["start"]) <= fend and e["kind"] in ("ramadan", "eid", "promotion", "adha_shopping", "eid_shopping")]
    return {
        "history": [{"t": d.strftime(fmt), "v": v} for d, v in hist],
        "forecast": [{"t": d.strftime(fmt), "v": v, "low": l, "high": h} for d, v, l, h in fore],
        "events": evs, "method": bt["best"], "error": bt["error"], "excluded_days": int(excluded.sum()),
        "table": [{"month": m["month"], "low": m["low"], "expected": m["expected"], "high": m["high"]} for m in fc["monthly"][:horizon]],
        "drivers": model.drivers(s, r["events"], fc),
        "options": _options(),
    }


def _options():
    d = service.load()["daily"]
    return {"products": sorted(d["product"].unique()), "categories": sorted(d["category"].unique()), "branches": sorted(d["branch"].unique())}


@app.get("/api/explorer/export.xlsx")
def explorer_export(product: str = "", category: str = "", branch: str = "", horizon: int = 3, user=Depends(current_user)):
    needs_data()
    from openpyxl import Workbook
    from openpyxl.styles import Font
    r = service.run(**_filters(product, category, branch))
    fc = r["forecast"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly forecast"
    sel = " / ".join(x for x in (product, category, branch) if x) or "Whole business"
    ws.append([f"Forecast: {sel}"]); ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Month", "Low (likely range)", "Expected", "High (likely range)"])
    for m in fc["monthly"][:horizon]:
        ws.append([m["month"], round(m["low"]), round(m["expected"]), round(m["high"])])
    ws2 = wb.create_sheet("Daily forecast")
    ws2.append(["Date", "Low", "Expected", "High"])
    end = pd.Timestamp(fc["monthly"][horizon - 1]["month"] + "-01") + pd.offsets.MonthEnd(0)
    for d, v in fc["daily"][fc["daily"].index <= end].items():
        ws2.append([d.date(), round(fc["daily_low"][d]), round(v), round(fc["daily_high"][d])])
    for w in (ws, ws2):
        for col in "ABCD":
            w.column_dimensions[col].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="ForecastIQ_forecast.xlsx"'})


@app.get("/api/breakdown")
def breakdown(by: str = "product", user=Depends(current_user)):
    needs_data()
    if by not in ("product", "category", "branch"):
        raise HTTPException(400, "by must be product, category or branch")
    return {"rows": service.breakdown(by)}


class ScenarioIn(BaseModel):
    discount: float = 0
    marketing: float = 0
    price: float = 0
    new_branch: bool = False
    branch_month: int = 1
    months: int = 3
    product: str = ""
    category: str = ""
    branch: str = ""
    name: str = ""


@app.post("/api/scenario")
def scenario(body: ScenarioIn, user=Depends(current_user)):
    needs_data()
    if not (0 <= body.discount <= 50 and -50 <= body.marketing <= 200 and 0 <= body.price <= 30 and 1 <= body.months <= 12):
        raise HTTPException(400, "Values out of range")
    return service.scenario(body.model_dump())


@app.get("/api/scenarios")
def scenarios(user=Depends(current_user)):
    rows = db.query("SELECT * FROM scenarios ORDER BY id DESC")
    for r in rows:
        r["params"] = json.loads(r.pop("params_json"))
    return rows


@app.post("/api/scenarios")
def save_scenario(body: ScenarioIn, user=Depends(EDIT)):
    res = service.scenario(body.model_dump())
    sid = db.execute("INSERT INTO scenarios(name, params_json, summary, created_by, created_at) VALUES(?,?,?,?,?)",
                     (body.name.strip() or "Scenario", json.dumps(body.model_dump()), res["summary"], user["name"], db.now()))
    return {"id": sid}


@app.delete("/api/scenarios/{sid}")
def delete_scenario(sid: int, user=Depends(EDIT)):
    db.execute("DELETE FROM scenarios WHERE id=?", (sid,))
    return {"ok": True}


@app.get("/api/accuracy")
def accuracy(user=Depends(current_user)):
    needs_data()
    return service.accuracy()


# --------------------------------------------------------------------------- data
@app.get("/api/datasets")
def datasets(user=Depends(current_user)):
    rows = db.query("SELECT id, name, filename, created_at, is_demo, active, health_json, decisions_json, uploaded_by FROM datasets ORDER BY id DESC")
    active = service.active_id()
    for r in rows:
        r["health"] = json.loads(r.pop("health_json") or "{}")
        r["decisions"] = json.loads(r.pop("decisions_json") or "{}")
        r["active"] = r["id"] == active
    return rows


@app.post("/api/datasets")
async def upload(file: UploadFile = File(...), stock: UploadFile | None = File(None), name: str = Form(""), user=Depends(EDIT)):
    content = await file.read()
    if len(content) > MAX_UPLOAD_MB * 1024 * 1024:
        raise HTTPException(400, f"Please keep the file under {MAX_UPLOAD_MB} MB")
    stock_bytes = await stock.read() if stock and stock.filename else None
    try:
        did = service.save_dataset(name.strip() or file.filename, file.filename, content, stock_bytes, user=user["name"])
    except data.DataError as e:
        raise HTTPException(400, str(e))
    db.execute("UPDATE datasets SET active=0")
    db.execute("UPDATE datasets SET active=1 WHERE id=?", (did,))
    service.invalidate()
    threading.Thread(target=service.warm, daemon=True).start()
    return {"id": did}


@app.post("/api/datasets/{did}/activate")
def activate(did: int, user=Depends(EDIT)):
    db.execute("UPDATE datasets SET active=0")
    db.execute("UPDATE datasets SET active=1 WHERE id=?", (did,))
    service.invalidate()
    threading.Thread(target=service.warm, daemon=True).start()
    return {"ok": True}


@app.post("/api/datasets/{did}/decision")
def decision(did: int, body: dict, user=Depends(EDIT)):
    d = db.one("SELECT decisions_json FROM datasets WHERE id=?", (did,))
    if not d:
        raise HTTPException(404, "Dataset not found")
    decs = json.loads(d["decisions_json"] or "{}")
    if body.get("key") == "missing" and body.get("value") in ("zero", "missing"):
        decs["missing"] = body["value"]
    service.redo_cleaning(did, decs)
    threading.Thread(target=service.warm, daemon=True).start()
    return {"ok": True}


@app.delete("/api/datasets/{did}")
def delete_dataset(did: int, user=Depends(need("admin"))):
    from .config import DATASETS_DIR
    for p in DATASETS_DIR.glob(f"{did}_*"):
        p.unlink(missing_ok=True)
    db.execute("DELETE FROM datasets WHERE id=?", (did,))
    service.invalidate()
    return {"ok": True}


@app.get("/api/template.csv")
def template():
    return FileResponse(SAMPLE_DIR / "sales_template.csv", filename="sales_template.csv", media_type="text/csv")


@app.get("/api/sample.csv")
def sample():
    return FileResponse(SAMPLE_DIR / "sahara_mart_sales.csv", filename="sample_sales_sahara_mart.csv", media_type="text/csv")


# --------------------------------------------------------------------------- reports
@app.get("/api/reports")
def list_reports(user=Depends(current_user)):
    from .config import SMTP_HOST
    return {"reports": db.query("SELECT id, month, title, sent_to, created_at FROM reports ORDER BY month DESC"), "email_connected": bool(SMTP_HOST)}


@app.post("/api/reports")
def make_report(user=Depends(EDIT)):
    needs_data()
    return {"id": reports.build()}


@app.get("/api/reports/{rid}/{fmt}")
def report_file(rid: int, fmt: str, user=Depends(current_user)):
    r = db.one("SELECT * FROM reports WHERE id=?", (rid,))
    if not r or fmt not in ("html", "pdf", "xlsx"):
        raise HTTPException(404, "Report not found")
    from pathlib import Path
    if fmt == "html":
        return HTMLResponse(Path(r["html_path"]).read_text())
    return FileResponse(r[f"{fmt}_path"], filename=f"ForecastIQ_report_{r['month']}.{fmt}")


@app.post("/api/reports/{rid}/send")
def send_report(rid: int, user=Depends(EDIT)):
    return {"sent_to": reports.send(rid)}


# --------------------------------------------------------------------------- settings, events, users
@app.get("/api/settings")
def get_settings(user=Depends(current_user)):
    return db.get_settings()


@app.put("/api/settings")
def put_settings(body: dict, user=Depends(need("admin"))):
    if "accent_color" in body and not re.fullmatch(r"#[0-9a-fA-F]{6}", str(body["accent_color"])):
        raise HTTPException(400, "Colour must look like #0d9488")
    if "report_recipients" in body:
        body["report_recipients"] = [e.strip() for e in body["report_recipients"] if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", e.strip())]
    db.set_settings(body)
    service.invalidate()
    return db.get_settings()


@app.post("/api/settings/logo")
async def upload_logo(file: UploadFile = File(...), user=Depends(need("admin"))):
    ext = (file.filename or "").rsplit(".", 1)[-1].lower()
    if ext not in ("png", "jpg", "jpeg", "svg", "webp"):
        raise HTTPException(400, "Please upload a PNG, JPG, SVG or WEBP logo")
    name = f"logo_{secrets.token_hex(4)}.{ext}"
    (BRAND_DIR / name).write_bytes(await file.read())
    db.set_settings({"logo_url": f"/brand/{name}"})
    return {"ok": True}


@app.get("/api/events")
def list_events(user=Depends(current_user)):
    return db.query("SELECT * FROM events ORDER BY start")


class EventIn(BaseModel):
    name: str
    start: str
    end: str
    kind: str = "promotion"
    discount: float = 0


@app.post("/api/events")
def add_event(body: EventIn, user=Depends(EDIT)):
    try:
        s, e = date.fromisoformat(body.start), date.fromisoformat(body.end)
    except ValueError:
        raise HTTPException(400, "Dates must look like 2026-11-11")
    if e < s or body.kind not in ("promotion", "holiday", "ramadan", "eid", "eid_shopping", "adha_shopping"):
        raise HTTPException(400, "Check the dates and type")
    db.execute("INSERT INTO events(name, start, end, kind, discount, source) VALUES(?,?,?,?,?,?)",
               (body.name.strip() or "Event", s.isoformat(), e.isoformat(), body.kind, max(0, min(body.discount, 90)) / 100 if body.discount > 1 else body.discount, "custom"))
    service.invalidate()
    threading.Thread(target=service.warm, daemon=True).start()
    return {"ok": True}


@app.delete("/api/events/{eid}")
def delete_event(eid: int, user=Depends(EDIT)):
    db.execute("DELETE FROM events WHERE id=?", (eid,))
    service.invalidate()
    return {"ok": True}


@app.get("/api/users")
def users(user=Depends(need("admin"))):
    return db.query("SELECT id, name, email, role FROM users ORDER BY id")


class UserIn(BaseModel):
    name: str
    email: str
    role: str = "viewer"


@app.post("/api/users")
def add_user(body: UserIn, user=Depends(need("admin"))):
    if body.role not in ("admin", "analyst", "viewer"):
        raise HTTPException(400, "Role must be admin, analyst or viewer")
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", body.email.strip()):
        raise HTTPException(400, "Please enter a valid email")
    if db.one("SELECT id FROM users WHERE lower(email)=lower(?)", (body.email.strip(),)):
        raise HTTPException(400, "This person already has access")
    temp = secrets.token_urlsafe(6)
    db.execute("INSERT INTO users(email, name, role, password_hash, created_at) VALUES(?,?,?,?,?)",
               (body.email.strip(), body.name.strip() or body.email, body.role, db.hash_password(temp), db.now()))
    return {"temporary_password": temp}


@app.patch("/api/users/{uid}")
def change_role(uid: int, body: dict, user=Depends(need("admin"))):
    if body.get("role") not in ("admin", "analyst", "viewer") or uid == user["id"]:
        raise HTTPException(400, "Can't change this role")
    db.execute("UPDATE users SET role=? WHERE id=?", (body["role"], uid))
    return {"ok": True}


@app.delete("/api/users/{uid}")
def remove_user(uid: int, user=Depends(need("admin"))):
    if uid == user["id"]:
        raise HTTPException(400, "You can't remove yourself")
    db.execute("DELETE FROM sessions WHERE user_id=?", (uid,))
    db.execute("DELETE FROM users WHERE id=?", (uid,))
    return {"ok": True}


# --------------------------------------------------------------------------- pages
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
app.mount("/brand", StaticFiles(directory=BRAND_DIR), name="brand")


@app.get("/health")
def health():
    return {"ok": True, "ready": STATE["ready"], "status": STATE["status"]}


@app.get("/landing")
def landing():
    return FileResponse(STATIC_DIR / "landing.html")


@app.get("/privacy")
def privacy():
    return FileResponse(STATIC_DIR / "privacy.html")


@app.get("/")
def index():
    return FileResponse(STATIC_DIR / "index.html", headers={"Cache-Control": "no-cache"})


@app.exception_handler(404)
async def not_found(request: Request, exc):
    if request.url.path.startswith("/api/"):
        return JSONResponse({"detail": getattr(exc, "detail", "Not found")}, status_code=404)
    return RedirectResponse("/")

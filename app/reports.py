"""Monthly Forecast Report: PDF (with chart), Excel, and a short email."""
from __future__ import annotations

from html import escape

import pandas as pd

from . import db, service
from .config import PUBLIC_BASE_URL, REPORTS_DIR


def _m(v):
    return f"Rs. {v / 1e6:,.2f}M" if abs(v) >= 1e6 else f"Rs. {v / 1e3:,.0f}k"


def build(months: int = 3) -> int:
    s = db.get_settings()
    ov = service.overview()
    r = service.run()
    fc = r["forecast"]["monthly"][:months]
    act = service.model.monthly_actual(r["series"])
    last12 = act[act.index <= pd.Timestamp(service._last_month(r["series"])[0])].iloc[-12:]
    prods = service.breakdown("product")
    branches = service.breakdown("branch")
    c = ov["cards"]
    month = fc[0]["month"]
    title = f"{pd.Timestamp(month + '-01'):%B} sales outlook: {_m(c['next_value'])} expected"
    base = REPORTS_DIR / f"forecast_{month}"

    # ---- PDF
    from reportlab.graphics.charts.lineplots import LinePlot
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.graphics.widgets.markers import makeMarker
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    accent = colors.HexColor(s["accent_color"])
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontSize=18, textColor=accent, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12.5, spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=9.5, leading=13)
    small = ParagraphStyle("s", parent=body, fontSize=8.5, textColor=colors.HexColor("#6b7280"))
    doc = SimpleDocTemplate(str(base.with_suffix(".pdf")), pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm)
    story = [Paragraph(escape(title), h1), Paragraph(f"{escape(s['company_name'])} · Monthly forecast report · data up to {ov['data_to']}", small), Spacer(1, 8)]
    nums = [(_m(c["last_month_value"]), f"{c['last_month']} actual"), (_m(c["next_value"]), f"{c['next_month']} forecast ({c['growth']:+.1f}%)"),
            (f"Rs. {c['low'] / 1e6:.2f}–{c['high'] / 1e6:.2f}M", "likely range (80%)"), (f"{c['accuracy']}%", "accuracy on recent months")]
    nt = Table([[Paragraph(f"<b><font size=13>{escape(v)}</font></b><br/><font size=7.5 color='#6b7280'>{escape(l)}</font>", body) for v, l in nums]],
               colWidths=[44.5 * mm] * 4, hAlign="LEFT")
    nt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f6f7f9")), ("INNERGRID", (0, 0), (-1, -1), 4, colors.white),
                            ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story.append(nt)
    # chart: last 12 months actual + forecast with range
    dw = Drawing(178 * mm, 62 * mm)
    lp = LinePlot()
    lp.x, lp.y, lp.width, lp.height = 14 * mm, 8 * mm, 160 * mm, 48 * mm
    xs_hist = list(range(len(last12)))
    xs_fc = list(range(len(last12) - 1, len(last12) + len(fc)))
    hist_pts = [(i, v / 1e6) for i, v in zip(xs_hist, last12.values)]
    first_fc = (len(last12) - 1, last12.values[-1] / 1e6)
    fc_pts = [first_fc] + [(len(last12) + i, m["expected"] / 1e6) for i, m in enumerate(fc)]
    lo_pts = [first_fc] + [(len(last12) + i, m["low"] / 1e6) for i, m in enumerate(fc)]
    hi_pts = [first_fc] + [(len(last12) + i, m["high"] / 1e6) for i, m in enumerate(fc)]
    lp.data = [hist_pts, fc_pts, lo_pts, hi_pts]
    lp.lines[0].strokeColor, lp.lines[0].strokeWidth = colors.HexColor("#111827"), 2
    lp.lines[1].strokeColor, lp.lines[1].strokeWidth, lp.lines[1].strokeDashArray = accent, 2, [4, 3]
    for i in (2, 3):
        lp.lines[i].strokeColor, lp.lines[i].strokeWidth, lp.lines[i].strokeDashArray = colors.HexColor("#99d5cf"), 1, [1, 2]
    lp.lines[1].symbol = makeMarker("FilledCircle", size=3)
    lp.xValueAxis.valueSteps = list(range(len(last12) + len(fc)))
    labels = [d.strftime("%b") for d in last12.index] + [pd.Timestamp(m["month"] + "-01").strftime("%b") for m in fc]
    lp.xValueAxis.labelTextFormat = lambda v: labels[int(v)] if 0 <= int(v) < len(labels) else ""
    lp.xValueAxis.labels.fontSize = lp.yValueAxis.labels.fontSize = 7
    lp.yValueAxis.labelTextFormat = "%.1fM"
    dw.add(lp)
    dw.add(String(14 * mm, 58 * mm, "Monthly sales: solid = actual, dashed = forecast, dotted = likely range", fontSize=7.5, fillColor=colors.HexColor("#6b7280")))
    story += [Spacer(1, 6), dw]
    story.append(Paragraph("What is driving the forecast", h2))
    story += [Paragraph("• " + escape(d), body) for d in ov["drivers"]]
    if ov["attention"]:
        story.append(Paragraph("Risks and opportunities", h2))
        story += [Paragraph("• " + escape(a["text"]), body) for a in ov["attention"]]
    story.append(Paragraph(f"Next {months} months", h2))
    t = Table([["Month", "Low", "Expected", "High"]] + [[pd.Timestamp(m["month"] + "-01").strftime("%B %Y"), _m(m["low"]), _m(m["expected"]), _m(m["high"])] for m in fc],
              hAlign="LEFT", colWidths=[45 * mm, 35 * mm, 35 * mm, 35 * mm])
    style = [("BACKGROUND", (0, 0), (-1, 0), accent), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
             ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e5e7eb")),
             ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f7f9")])]
    t.setStyle(TableStyle(style))
    story.append(t)
    story.append(Paragraph("Products", h2))
    pt = Table([["Product", "Last month", "Next month", "Change", "Stock"]] +
               [[p["name"], _m(p["last_month"]), _m(p["next_month"]), f"{p['growth']:+.1f}%" if p["growth"] is not None else "—", p["risk"] or "—"] for p in prods],
               hAlign="LEFT", colWidths=[62 * mm, 28 * mm, 28 * mm, 20 * mm, 32 * mm])
    pt.setStyle(TableStyle(style))
    story.append(pt)
    story.append(Paragraph("Branches", h2))
    bt = Table([["Branch", "Last month", "Next month", "Change"]] + [[b["name"], _m(b["last_month"]), _m(b["next_month"]), f"{b['growth']:+.1f}%"] for b in branches],
               hAlign="LEFT", colWidths=[62 * mm, 28 * mm, 28 * mm, 20 * mm])
    bt.setStyle(TableStyle(style))
    story += [bt, Spacer(1, 8), Paragraph("Forecasts are estimates with ranges, not promises. The likely range covers about 8 in 10 outcomes based on how accurate "
                                          "the forecast was on your recent months.", small)]
    doc.build(story)

    # ---- Excel
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"
    ws.append(["Month", "Low", "Expected", "High"])
    for m in r["forecast"]["monthly"]:
        ws.append([m["month"], round(m["low"]), round(m["expected"]), round(m["high"])])
    ws2 = wb.create_sheet("Products")
    ws2.append(["Product", "Last month", "Next month forecast", "Change %", "Stock risk", "Backtest error %"])
    for p in prods:
        ws2.append([p["name"], round(p["last_month"]), round(p["next_month"]), round(p["growth"], 1) if p["growth"] is not None else None, p["risk"],
                    round(p["error"] * 100, 1) if p["error"] is not None else None])
    ws3 = wb.create_sheet("Branches")
    ws3.append(["Branch", "Last month", "Next month forecast", "Change %"])
    for b in branches:
        ws3.append([b["name"], round(b["last_month"]), round(b["next_month"]), round(b["growth"], 1)])
    for w in (ws, ws2, ws3):
        for cell in w[1]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=s["accent_color"].lstrip("#"))
        for col in "ABCDEF":
            w.column_dimensions[col].width = 22
    wb.save(base.with_suffix(".xlsx"))

    # ---- email
    link = f"{PUBLIC_BASE_URL}/#/" if PUBLIC_BASE_URL else "#"
    bullets = "".join(f"<li style='margin-bottom:6px'>{escape(d)}</li>" for d in ov["drivers"][:3])
    html = f"""<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{escape(title)}</title></head>
<body style="margin:0;background:#f3f4f6;font-family:Arial,Helvetica,sans-serif;color:#111827"><table role="presentation" width="100%"><tr><td align="center" style="padding:16px 8px">
<table role="presentation" width="100%" style="max-width:560px;background:#fff;border-radius:14px;overflow:hidden">
<tr><td style="background:{s['accent_color']};color:#fff;padding:18px 20px"><div style="font-size:13px;opacity:.85">{escape(s['company_name'])} · ForecastIQ</div>
<div style="font-size:21px;font-weight:700;margin-top:4px">{escape(title)}</div></td></tr>
<tr><td style="padding:14px"><table role="presentation" width="100%"><tr>
{''.join(f'<td style="padding:5px;width:33%"><div style="background:#f6f7f9;border-radius:10px;padding:12px 6px;text-align:center"><div style="font-size:19px;font-weight:800">{escape(v)}</div><div style="font-size:11.5px;color:#6b7280">{escape(l)}</div></div></td>'
          for v, l in [(_m(c['next_value']), 'expected'), (f"{c['low'] / 1e6:.2f}–{c['high'] / 1e6:.2f}M", 'likely range'), (f"{c['growth']:+.1f}%", 'vs ' + c['last_month'].split()[0])])}
</tr></table></td></tr>
<tr><td style="padding:0 20px"><h3 style="font-size:15px;margin:8px 0 6px">Why</h3><ul style="padding-left:18px;font-size:14px;margin:0">{bullets}</ul></td></tr>
<tr><td align="center" style="padding:18px 20px 22px"><a href="{link}" style="background:{s['accent_color']};color:#fff;text-decoration:none;padding:12px 22px;border-radius:9px;font-weight:700;display:inline-block">Open dashboard</a>
<div style="font-size:12px;color:#6b7280;margin-top:10px">Full report attached as PDF and Excel.</div></td></tr></table></td></tr></table></body></html>"""
    base.with_suffix(".html").write_text(html)
    existing = db.one("SELECT id FROM reports WHERE month=?", (month,))
    if existing:
        db.execute("UPDATE reports SET title=?, html_path=?, pdf_path=?, xlsx_path=?, created_at=? WHERE id=?",
                   (title, str(base.with_suffix(".html")), str(base.with_suffix(".pdf")), str(base.with_suffix(".xlsx")), db.now(), existing["id"]))
        return existing["id"]
    return db.execute("INSERT INTO reports(month, title, html_path, pdf_path, xlsx_path, created_at) VALUES(?,?,?,?,?,?)",
                      (month, title, str(base.with_suffix(".html")), str(base.with_suffix(".pdf")), str(base.with_suffix(".xlsx")), db.now()))


def send(rid: int) -> str:
    import smtplib
    from email.message import EmailMessage
    from pathlib import Path
    from .config import SMTP_FROM, SMTP_HOST, SMTP_PASSWORD, SMTP_PORT, SMTP_USER
    r = db.one("SELECT * FROM reports WHERE id=?", (rid,))
    to = db.get_settings()["report_recipients"]
    if not SMTP_HOST or not to:
        db.execute("UPDATE reports SET sent_to='outbox' WHERE id=?", (rid,))
        return "outbox"
    msg = EmailMessage()
    msg["Subject"], msg["From"], msg["To"] = r["title"], SMTP_FROM, ", ".join(to)
    msg.set_content(r["title"])
    msg.add_alternative(Path(r["html_path"]).read_text(), subtype="html")
    for p, sub in ((r["pdf_path"], "pdf"), (r["xlsx_path"], "vnd.openxmlformats-officedocument.spreadsheetml.sheet")):
        msg.add_attachment(Path(p).read_bytes(), maintype="application", subtype=sub, filename=Path(p).name)
    try:
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=30) as smtp:
            smtp.starttls()
            if SMTP_USER:
                smtp.login(SMTP_USER, SMTP_PASSWORD)
            smtp.send_message(msg)
        db.execute("UPDATE reports SET sent_to=? WHERE id=?", (", ".join(to), rid))
        return ", ".join(to)
    except Exception:
        db.execute("UPDATE reports SET sent_to='outbox (send failed)' WHERE id=?", (rid,))
        return "outbox"
